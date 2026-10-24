from pathlib import Path
import re

from pypdf import PdfReader
from korsik_pdf_ocr import pdf_pages_with_ocr
from docx import Document
from pptx import Presentation


DOCS_DIR = Path(__file__).resolve().parent / "data" / "knowledge_docs"
EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def read_document(path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        yield from pdf_pages_with_ocr(path)

    elif suffix == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        yield "документ", "\n".join(parts)

    elif suffix == ".pptx":
        presentation = Presentation(str(path))
        for number, slide in enumerate(presentation.slides, 1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
            yield f"слайд {number}", "\n".join(parts)

    elif suffix in {".txt", ".md"}:
        yield "документ", path.read_text(encoding="utf-8")


def split_text(text, size=1100, overlap=150):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return

    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            space = text.rfind(" ", start + size // 2, end)
            if space > start:
                end = space

        chunk = text[start:end].strip()
        if chunk:
            yield chunk

        if end >= len(text):
            break

        start = max(start + 1, end - overlap)


def load_documents():
    fragments = []

    if not DOCS_DIR.exists():
        return fragments

    for path in sorted(
        list(DOCS_DIR.rglob("*"))
        + list((DOCS_DIR.parent / "policy_docs").rglob("*"))
    ):
        if not path.is_file() or path.suffix.lower() not in EXTENSIONS:
            continue

        try:
            for location, text in read_document(path):
                for chunk in split_text(text):
                    fragments.append({
                        "source": path.name,
                        "location": location,
                        "text": chunk,
                    })
        except Exception as exc:
            print(f"Не удалось прочитать {path.name}: {exc}")

    return fragments


def search_documents(question, limit=4):
    words = set(re.findall(r"[а-яёa-z0-9]{3,}", question.lower()))
    if not words:
        return []

    scored = []

    for fragment in load_documents():
        text = fragment["text"].lower()
        score = sum(text.count(word) for word in words)

        if score > 0:
            scored.append((score, fragment))

    scored.sort(key=lambda item: item[0], reverse=True)
    return [fragment for _, fragment in scored[:limit]]
