"""Импорт сценария Корсика из Excel в SQLite и файл для работающего онбординга."""

import sys as _korsik_sys
from pathlib import Path as _korsik_Path
_korsik_root = _korsik_Path(__file__).resolve().parent.parent
if str(_korsik_root) not in _korsik_sys.path:
    _korsik_sys.path.insert(0, str(_korsik_root))
import json
import re
import sys
from pathlib import Path
from openpyxl import load_workbook
from database import connect, init_db

ROOT = Path(__file__).resolve().parent.parent
HEADERS = ('Месяц/день/время информирования', 'Информационный повод', 'Тег сотрудника', 'Форма/Макет', 'Текст сообщения', 'Обязательный этап', 'Сопроводительное изображение')
VALID_TAGS = {'gtp': 'GTP', 'grpp': 'GrPP', 'luzhniki': 'Luzhniki'}

def clean(value):
    if value is None: return ''
    return str(value).strip()

def read_steps(path, sheet_name=None):
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book[sheet_name] if sheet_name else next((s for s in book if any('Текст сообщения' == clean(v) for v in next(s.iter_rows(min_row=2, max_row=2, values_only=True)))), book.active)
    headers = [clean(x) for x in next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))]
    missing = [h for h in HEADERS if h not in headers]
    if missing: raise ValueError('В Excel отсутствуют колонки: ' + ', '.join(missing))
    result = []
    for row_num, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        row = dict(zip(headers, values))
        message = clean(row.get('Текст сообщения'))
        if not message: continue
        tag_text = clean(row.get('Тег сотрудника'))
        tags = []
        for word in re.split(r'[,;/\n]+', tag_text):
            tag = VALID_TAGS.get(word.strip().casefold())
            if tag and tag not in tags: tags.append(tag)
                                                                                                        
        options = ['Далее']
        if row_num == 5: options = ['На ТЫ', 'На ВЫ']
        if row_num == 25: options = ['Да!', 'Нет, спасибо!']
                                                                                        
        result.append({
            'id': f'excel-{row_num}', 'source_row': row_num,
            'time': clean(row.get('Месяц/день/время информирования')),
            'reason': clean(row.get('Информационный повод')),
            'tags': tags, 'form': clean(row.get('Форма/Макет')),
            'text': message, 'buttons': options,
            'mandatory': clean(row.get('Обязательный этап')).casefold() in ('да', 'yes', 'true', '1'),
            'attachment': clean(row.get('Сопроводительное изображение')),
            'conditional': row_num in (26, 27, 28),
        })
    return sheet.title, result

def import_excel(path, sheet_name=None, output=None):
    path = Path(path)
    if not path.is_file(): raise FileNotFoundError(path)
    sheet, steps = read_steps(path, sheet_name)
    if not steps: raise ValueError('В Excel не найдено ни одного сообщения')
    init_db()
    schema = (ROOT / 'schema.sql').read_text(encoding='utf-8')
                                                                   
    with connect() as conn:
        conn.executescript(schema)
    with connect() as conn:
        conn.execute('BEGIN IMMEDIATE')
        conn.execute('INSERT OR IGNORE INTO onboarding_templates(name,version,description,is_active) VALUES(?,?,?,1)', ('базовый сценарий адаптации', 1, f'Импорт из {path.name}'))
        template_id = conn.execute('SELECT id FROM onboarding_templates WHERE name=? AND version=1', ('базовый сценарий адаптации',)).fetchone()['id']
                                                                                          
        for order, step in enumerate(steps, 1):
            conn.execute('''INSERT INTO scenario_steps(template_id,step_key,source_row,step_order,schedule_text,trigger_type,information_reason,employee_tags,form_name,message_text,mandatory,attachment,extra_notes)
              VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
              ON CONFLICT(template_id,step_key) DO UPDATE SET source_row=excluded.source_row,step_order=excluded.step_order,schedule_text=excluded.schedule_text,trigger_type=excluded.trigger_type,information_reason=excluded.information_reason,employee_tags=excluded.employee_tags,form_name=excluded.form_name,message_text=excluded.message_text,mandatory=excluded.mandatory,attachment=excluded.attachment,extra_notes=excluded.extra_notes''',
              (template_id,step['id'],step['source_row'],order,step['time'],'scheduled' if step['time'] else 'manual',step['reason'],json.dumps(step['tags'],ensure_ascii=False),step['form'],step['text'],int(step['mandatory']),step['attachment'],json.dumps({'buttons':step['buttons'],'conditional':step['conditional']},ensure_ascii=False)))
                                                                                            
    destination = Path(output) if output else ROOT / 'data' / 'corsik_scenario.json'
    destination.parent.mkdir(parents=True,exist_ok=True)
    temp = destination.with_suffix('.json.tmp')
    temp.write_text(json.dumps(steps,ensure_ascii=False,indent=2),encoding='utf-8')
    temp.replace(destination)
    return {'sheet':sheet,'imported':len(steps),'tagged':sum(bool(s['tags']) for s in steps),'output':str(destination)}

if __name__ == '__main__':
    if len(sys.argv)<2: raise SystemExit('Использование: python import_excel.py путь_к_excel.xlsx')
    print(import_excel(sys.argv[1]))
